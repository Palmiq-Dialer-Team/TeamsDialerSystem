"""Save disposition results to Excel files."""

import threading
import os
import json
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, Optional
import openpyxl
from openpyxl import Workbook, load_workbook
import time
import pythoncom
import win32com.client
import win32gui
import pandas as pd

# Excel window state constant (restore from minimized)
XL_NORMAL = -4143  # xlNormal


class DispositionSaver:
    """Saves dispositions to Excel files."""
    
    def __init__(self, base_folder: str, active_workbook=None, workbook_path=None, colors=None):
        self.base_folder = Path(base_folder)
        self.dispositions_folder = self.base_folder / "Dispositions"
        self.max_retries = 3
        self.retry_delay = 1.0  # seconds
        self.active_workbook = active_workbook
        self.workbook_path = workbook_path
        default_colors = {
            "Booked Meeting": "#C6EFCE",      # Pastel Green
            "Sent Info": "#B3E5FC",           # Pastel Blue
            "Call Back Later": "#FFEB9C",     # Pastel Yellow-Orange
            "Voicemail": "#E1BEE7",           # Pastel Purple (auto-detected)
            "Voicemail (Verified)": "#E1BEE7",# Pastel Purple (manual)
            "No Response": "#F2F2F2",         # Very Light Gray
            "Not Interested": "#FFCCBC",      # Pastel Peach
            "Wrong Number": "#FFCDD2",        # Pastel Red
            "Invalid Contact": "#EF9A9A"      # Pastel Darker Red
        }
        self.disposition_colors = colors if colors else default_colors
        # Dispositions for which the active phone number cell is bolded in source + disposition files
        self.dispositions_bold_active_phone = frozenset({
            "Voicemail (Verified)", "Booked Meeting", "Not Interested", "Sent Info", "Call Back Later"
        })
    
    def update_colors(self, new_colors):
        """Update disposition colors."""
        self.disposition_colors = new_colors
    
    def hex_to_rgb_int(self, hex_color):
        """Hex to Excel BGR int."""
        h = hex_color.lstrip('#')
        r, g, b = tuple(int(h[i:i+2], 16) for i in (0, 2, 4))
        return r + (g * 256) + (b * 256 * 256)

    def hex_to_rgb_long(self, hex_color):
        """Hex to Excel RGB Long for AutoFilter."""
        h = hex_color.lstrip('#')
        r, g, b = tuple(int(h[i:i+2], 16) for i in (0, 2, 4))
        return r + (g * 256) + (b * 65536)
    
    def _get_disposition_file_path(self, disposition_type: str) -> Path:
        filename = f"{disposition_type}.xlsx"
        return self.dispositions_folder / filename
    
    def _ensure_folder_exists(self, folder_path: Path):
        folder_path.mkdir(parents=True, exist_ok=True)
    
    def _copy_row_from_source(self, workbook, sheet_name: str, row_number: int) -> list:
        try:
            if sheet_name:
                try:
                    worksheet = workbook.Worksheets(sheet_name)
                except:
                    worksheet = workbook.ActiveSheet
            else:
                worksheet = workbook.ActiveSheet
            
            if worksheet.UsedRange:
                max_col = worksheet.UsedRange.Columns.Count
            else:
                max_col = 0
            
            row_data = []
            for col_idx in range(1, max_col + 1):
                cell = worksheet.Cells(row_number, col_idx)
                value = cell.Value
                row_data.append(value)
            
            return row_data
        except Exception as e:
            print(f"Error copying row from source: {e}")
            return []
    
    def _get_source_headers(self, workbook, sheet_name: str) -> list:
        return self._copy_row_from_source(workbook, sheet_name, 1)
    
    def _create_new_file(self, file_path: Path, headers: list, row_data: list, bold_phone_header: Optional[str] = None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        
        bold_col_idx = None
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            if bold_phone_header and header == bold_phone_header:
                from openpyxl.styles import Font
                cell.font = Font(bold=True)
                bold_col_idx = col_idx
        
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=2, column=col_idx, value=value)
            if bold_col_idx is not None and col_idx == bold_col_idx:
                from openpyxl.styles import Font
                cell.font = Font(bold=True)
        
        wb.save(file_path)
        print(f"Created new disposition file: {file_path}")
    
    def _append_to_existing_file(self, file_path: Path, headers: list, row_data: list, bold_phone_header: Optional[str] = None):
        wb = load_workbook(file_path)
        ws = wb.active
        
        existing_headers = []
        for cell in ws[1]:
            existing_headers.append(cell.value)
        
        next_row = ws.max_row + 1
        
        clean_row_data = []
        for col_idx, header in enumerate(existing_headers, start=1):
            try:
                source_idx = headers.index(header)
                if source_idx < len(row_data):
                    val = row_data[source_idx]
                else:
                    val = None
            except ValueError:
                val = None
            
            clean_row_data.append(val)
            
        bold_col_idx = None
        if bold_phone_header and bold_phone_header in existing_headers:
            bold_col_idx = existing_headers.index(bold_phone_header) + 1
        
        for col_idx, value in enumerate(clean_row_data, start=1):
            if isinstance(value, str):
                value = "".join(ch for ch in value if (0x20 <= ord(ch) <= 0xD7FF))
                
            cell = ws.cell(row=next_row, column=col_idx, value=value)
            if bold_col_idx is not None and col_idx == bold_col_idx:
                from openpyxl.styles import Font
                cell.font = Font(bold=True)
        
        wb.save(file_path)
        print(f"Appended row to disposition file: {file_path}")
    
    def _save_contact_sync(self, contact_data: Dict[str, Any], disposition: str):
        self._ensure_folder_exists(self.dispositions_folder)
        file_path = self._get_disposition_file_path(disposition)
        
        workbook = self._get_workbook_in_thread()
        if not workbook:
            print(f"Error: Could not access source workbook to copy row for disposition '{disposition}'")
            return
        
        try:
            row_number = contact_data.get('rowNumber')
            sheet_name = contact_data.get('sheetName')
            
            if not row_number:
                print(f"Error: No rowNumber in contact_data for disposition '{disposition}'")
                return
            
            row_data = self._copy_row_from_source(workbook, sheet_name, row_number)
            if not row_data:
                print(f"Error: Could not copy row {row_number} from source file for disposition '{disposition}'")
                return
            
            headers = self._get_source_headers(workbook, sheet_name)
            if not headers:
                print(f"Warning: Could not get headers from source file, using row data length")
                headers = [f"Column_{i+1}" for i in range(len(row_data))]
            
            if "Original Row Index" not in headers:
                headers.append("Original Row Index")
                row_data.append(row_number)  # Add the row number as the value
            else:
                idx = headers.index("Original Row Index")
                if idx < len(row_data):
                    row_data[idx] = row_number
                else:
                    row_data.append(row_number)
            
            # Determine which phone column (header) should be bolded in the disposition file
            bold_phone_header = None
            if disposition in self.dispositions_bold_active_phone:
                active_idx = contact_data.get("ActivePhoneIndex")
                phone_objs = contact_data.get("Phone_Objects") or []
                if isinstance(active_idx, int) and 0 <= active_idx < len(phone_objs):
                    bold_phone_header = phone_objs[active_idx].get("original_col")
                    # Ensure header list includes the column name
                    if bold_phone_header and bold_phone_header not in headers:
                        headers.append(bold_phone_header)
                        row_data.append(None)
        except Exception as e:
            print(f"Error copying row from source file: {e}")
            import traceback
            traceback.print_exc()
            return
        
        current_retry_delay = self.retry_delay
        for attempt in range(self.max_retries):
            try:
                if file_path.exists():
                    self._append_to_existing_file(file_path, headers, row_data, bold_phone_header=bold_phone_header)
                else:
                    self._create_new_file(file_path, headers, row_data, bold_phone_header=bold_phone_header)
                
                print(f"Successfully saved row {row_number} to '{disposition}.xlsx'")
                break
                
            except PermissionError as e:
                if attempt < self.max_retries - 1:
                    print(f"File locked, retrying in {current_retry_delay} seconds... (Attempt {attempt + 1}/{self.max_retries})")
                    time.sleep(current_retry_delay)
                    current_retry_delay *= 1.5
                else:
                    print(f"Error: Failed to save after {self.max_retries} attempts: {e}")
                    raise
            except Exception as e:
                print(f"Error saving contact data to '{disposition}.xlsx': {e}")
                import traceback
                traceback.print_exc()
                raise
    
    def _get_workbook_in_thread(self):
        """Re-acquire workbook COM in this thread. Retries when Excel is busy."""
        if not self.workbook_path:
            return None
        RPC_E_CALL_REJECTED = -2147418111
        max_retries = 3
        retry_delay = 0.4  # seconds
        for attempt in range(max_retries):
            try:
                excel_app = win32com.client.GetActiveObject("Excel.Application")
                for workbook in excel_app.Workbooks:
                    if workbook.FullName == self.workbook_path:
                        return workbook
                if excel_app.ActiveWorkbook:
                    return excel_app.ActiveWorkbook
                return None
            except Exception as e:
                err_code = getattr(e, 'args', [None])[0] if e.args else None
                if err_code == RPC_E_CALL_REJECTED and attempt < max_retries - 1:
                    time.sleep(retry_delay)
                    continue
                print(f"Error re-acquiring workbook: {e}")
                return None
        return None
    
    def _bring_excel_to_foreground(self, workbook):
        """Restore and bring the Excel window to the top so coloring/bolding apply when not minimized."""
        try:
            app = workbook.Application
            app.Visible = True
            app.WindowState = XL_NORMAL  # restore if minimized
            hwnd = int(app.Hwnd)
            win32gui.SetForegroundWindow(hwnd)
        except Exception as e:
            print(f"Could not bring Excel to front: {e}")

    def _update_source_file(self, contact_data: Dict[str, Any], disposition: str):
        """Color row and stamp _DISPO_STAMP. Call from COM-initialized thread."""
        workbook = self._get_workbook_in_thread()
        if not workbook:
            print("Warning: Could not access active workbook for live editing.")
            return

        self._bring_excel_to_foreground(workbook)
        
        try:
            target_sheet_name = str(contact_data.get('sheetName')).strip()
            worksheet = None

            for s in workbook.Worksheets:
                if s.Name.strip().lower() == target_sheet_name.lower():
                    worksheet = s
                    break
            
            if not worksheet:
                print(f"Could not find sheet '{target_sheet_name}'.")
                worksheet = workbook.ActiveSheet 

            row_number = contact_data.get('rowNumber')
            
            if not row_number or int(row_number) < 1 or int(row_number) > 1048576:
                print(f"Invalid row number {row_number}")
                return
            
            # Get the row object
            row = worksheet.Rows(row_number)
            
            # Color entire row based on disposition
            hex_color = self.disposition_colors.get(disposition, "#FFFFFF")
            row.Interior.Color = self.hex_to_rgb_int(hex_color)

            # Bold the active phone number cell in the source sheet for Voicemail (Verified), Booked Meeting, Not Interested, Sent Info, Call Back Later
            if disposition in self.dispositions_bold_active_phone:
                active_idx = contact_data.get("ActivePhoneIndex")
                phone_objs = contact_data.get("Phone_Objects") or []
                if isinstance(active_idx, int) and 0 <= active_idx < len(phone_objs):
                    phone_header = phone_objs[active_idx].get("original_col")
                    if phone_header:
                        max_col = worksheet.UsedRange.Columns.Count if worksheet.UsedRange else 0
                        phone_col = None
                        for col_idx in range(1, max_col + 1):
                            cell_value = worksheet.Cells(1, col_idx).Value
                            if cell_value and str(cell_value).strip() == str(phone_header).strip():
                                phone_col = col_idx
                                break
                        if phone_col:
                            phone_cell = worksheet.Cells(row_number, phone_col)
                            try:
                                phone_cell.Font.Bold = True
                            except Exception:
                                pass
            
            stamp_column = None
            header_row = worksheet.Rows(1)
            max_col = worksheet.UsedRange.Columns.Count if worksheet.UsedRange else 0
            
            # Search existing columns
            for col_idx in range(1, max_col + 1):
                cell_value = worksheet.Cells(1, col_idx).Value
                if cell_value and str(cell_value).strip().upper() == "_DISPO_STAMP":
                    stamp_column = col_idx
                    break
            
            if stamp_column is None:
                stamp_column = max_col + 1
                worksheet.Cells(1, stamp_column).Value = "_DISPO_STAMP"
                header_cell = worksheet.Cells(1, stamp_column)
                header_cell.Font.Bold = True
                header_cell.Interior.ColorIndex = 15
            
            stamp_data = {
                "disposition": disposition,
                "timestamp": datetime.now().isoformat(),
                "notes": contact_data.get('Notes', '')
            }
            stamp_json = json.dumps(stamp_data)
            
            worksheet.Cells(row_number, stamp_column).Value = stamp_json
            
            notes_col = None
            for col_idx in range(1, max_col + 1):
                cell_value = worksheet.Cells(1, col_idx).Value
                if cell_value and str(cell_value).strip().lower() == "notes":
                    notes_col = col_idx
                    break
            
            if notes_col:
                notes_value = contact_data.get('Notes', '')
                worksheet.Cells(row_number, notes_col).Value = notes_value
            
            workbook.Save()
            
            print(f"Updated source file: Row {row_number} colored and stamped with '{disposition}'")
            
        except Exception as e:
            print(f"Error updating source file: {e}")
            import traceback
            traceback.print_exc()
    
    def save_contact_async(self, contact_data: Dict[str, Any], disposition: str):
        """Save contact in background: update source, then save to disposition files."""
        def save_thread():
            pythoncom.CoInitialize()
            try:
                self._update_source_file(contact_data, disposition)
                self._save_contact_sync(contact_data, disposition)
            except Exception as e:
                print(f"Error in async save thread: {e}")
                import traceback
                traceback.print_exc()
            finally:
                pythoncom.CoUninitialize()
        
        thread = threading.Thread(target=save_thread, daemon=True)
        thread.start()

    def save_workbook_async(self):
        """Save active workbook in background (autosave)."""
        if not self.workbook_path:
            return
        def save_thread():
            pythoncom.CoInitialize()
            try:
                workbook = self._get_workbook_in_thread()
                if workbook and hasattr(workbook, 'Saved') and not workbook.Saved:
                    workbook.Save()
                    return True
            except Exception as e:
                print(f"Autosave error: {e}")
            finally:
                pythoncom.CoUninitialize()
            return False
        thread = threading.Thread(target=save_thread, daemon=True)
        thread.start()
    
    def update_source_file(self, workbook_path: str, sheet_name: str, row_index: int, disposition: str, notes: str = None):
        """
        Update the original source file with the disposition, using a closed-file edit only.

        - For .xlsx/.xlsm: use openpyxl.
        - For .csv: update the Disposition Stamp (and Notes if present) via pandas.

        Args:
            workbook_path: Path to the workbook or CSV
            sheet_name: Name of the sheet (ignored for CSV)
            row_index: Row index (0-based from pandas data, or Excel rowNumber-2 for CSV)
            disposition: Disposition type
            notes: Optional notes to save to the Notes column
        """
        from pathlib import Path
        path = Path(workbook_path)
        suffix = path.suffix.lower()

        # --- CSV path: update only Disposition Stamp (and Notes if present) ---
        if suffix == ".csv":
            try:
                import pandas as pd
                df = pd.read_csv(path)
            except Exception as e:
                print(f"CSV update failed: could not read {path}: {e}")
                return

            # row_index is expected to be a 0-based data index for CSV
            idx = row_index
            if not isinstance(idx, int) or idx < 0 or idx >= len(df):
                print(f"CSV update skipped: row index {idx} out of range for {path}")
                return

            if "Disposition Stamp" in df.columns:
                df.loc[idx, "Disposition Stamp"] = disposition

            if notes is not None and "Notes" in df.columns:
                df.loc[idx, "Notes"] = notes

            try:
                df.to_csv(path, index=False)
                print(f"Updated CSV source file: row index {idx} in {path}")
            except Exception as e:
                print(f"Failed to write CSV {path}: {e}")
            return

        # --- XLSX/xlsm path: closed-file editing with openpyxl ---
        try:
            from openpyxl.styles import PatternFill

            # Use dynamic colors from settings, with fallback color_map
            hex_color = self.disposition_colors.get(disposition)
            if not hex_color:
                color_map = {
                    "Booked Meeting": "C6EFCE",
                    "Sent Info": "B3E5FC",
                    "Call Back Later": "FFEB9C",
                    "Voicemail (Verified)": "E1BEE7",
                    "No Response": "F2F2F2",
                    "Not Interested": "FFCCBC",
                    "Wrong Number": "FFCDD2",
                    "Invalid Contact": "EF9A9A"
                }
                clean_hex = color_map.get(disposition, "FFFFFF")  # Default to white if unknown
            else:
                clean_hex = hex_color.lstrip('#')  # openpyxl wants 'RRGGBB'

            color_fill = PatternFill(start_color=clean_hex, end_color=clean_hex, fill_type="solid")

            workbook = load_workbook(workbook_path)

            if sheet_name not in workbook.sheetnames:
                print(f"ERROR: Sheet '{sheet_name}' not found in workbook '{workbook_path}'.")
                return

            sheet = workbook[sheet_name]

            # Find the columns for "Disposition Stamp", "Status Color", and "Notes"
            disposition_col = None
            color_col = None
            notes_col = None
            for col in range(1, sheet.max_column + 1):
                header = sheet.cell(row=1, column=col).value
                if header and str(header).strip():
                    header_str = str(header).strip().lower()
                    if header_str == "disposition stamp":
                        disposition_col = col
                    elif header_str == "status color":
                        color_col = col
                    elif header_str == "notes":
                        notes_col = col

            if not disposition_col or not color_col:
                print("ERROR: 'Disposition Stamp' or 'Status Color' columns not found in source file.")
                return

            # openpyxl uses 1-based indexing; row_index is 0-based data index => add 2 (1 header, 1 for 0-base)
            target_row = row_index + 2

            sheet.cell(row=target_row, column=disposition_col).value = disposition
            sheet.cell(row=target_row, column=color_col).fill = color_fill

            if notes_col and notes is not None:
                sheet.cell(row=target_row, column=notes_col).value = notes

            workbook.save(workbook_path)
            print(f"Successfully updated closed source file: {workbook_path} (row index {row_index})")

        except PermissionError:
            print(f"FATAL ERROR: Could not update source file with openpyxl. The file is likely open and locked by Excel. Please close the file and try again.")
        except Exception as ex_openpyxl:
            print(f"ERROR: An unexpected error occurred with openpyxl: {ex_openpyxl}")
            import traceback
            traceback.print_exc()
    
    def remove_contact_from_file(self, contact_data: dict, disposition_name: str):
        """
        Removes a contact's row from a specific disposition Excel file.

        Args:
            contact_data: The dictionary of the contact to remove.
            disposition_name: The name of the disposition file (e.g., "Call Back Later").
        """
        disposition_filename = f"{disposition_name}.xlsx"
        disposition_filepath = os.path.join(self.base_folder, "Dispositions", disposition_filename)

        if not os.path.exists(disposition_filepath):
            print(f"File not found, cannot remove contact: {disposition_filepath}")
            return

        try:
            df = pd.read_excel(disposition_filepath)

            # Identify the row to remove using 'Original Row Index' as a unique key
            # Fallback to 'rowNumber' if 'Original Row Index' is not available
            original_row_idx = contact_data.get('Original Row Index')
            if original_row_idx is None:
                original_row_idx = contact_data.get('rowNumber')
            
            if original_row_idx is None:
                print("Cannot remove contact: 'Original Row Index' or 'rowNumber' not found in contact data.")
                return

            # Check if the column exists in the DataFrame
            if 'Original Row Index' not in df.columns and 'rowNumber' not in df.columns:
                print("Cannot remove contact: 'Original Row Index' or 'rowNumber' column not found in file.")
                return

            # Use the appropriate column name
            column_name = 'Original Row Index' if 'Original Row Index' in df.columns else 'rowNumber'

            # Keep all rows that DO NOT match the original row index
            initial_rows = len(df)
            df_cleaned = df[df[column_name] != original_row_idx]
            final_rows = len(df_cleaned)

            if initial_rows == final_rows:
                print(f"Contact with {column_name} '{original_row_idx}' not found in '{disposition_filename}'.")
            else:
                # Save the cleaned DataFrame back to the Excel file
                df_cleaned.to_excel(disposition_filepath, index=False)
                print(f"Successfully removed contact from '{disposition_filename}'.")

        except Exception as e:
            print(f"ERROR: Failed to remove contact from {disposition_filename}: {e}")
            import traceback
            traceback.print_exc()
    
    def revert_source_file_changes(self, contact_data: Dict[str, Any]):
        """
        Reverts the disposition changes in the source file (clears color and stamp).
        Handles both live editing and closed file editing.
        
        Args:
            contact_data: Dictionary containing contact data (must include 'rowNumber' and 'sheetName')
        """
        try:
            # --- ATTEMPT 1: LIVE EDITING ---
            workbook = self._get_workbook_in_thread()
            if workbook:
                # Get the worksheet name
                sheet_name = contact_data.get('sheetName')
                if not sheet_name or not isinstance(sheet_name, str) or len(sheet_name.strip()) == 0:
                    sheet_name = None
                else:
                    sheet_name = str(sheet_name).strip()
                
                # Get the worksheet
                if sheet_name:
                    try:
                        worksheet = workbook.Worksheets(sheet_name)
                    except Exception as e:
                        worksheet = workbook.ActiveSheet
                else:
                    worksheet = workbook.ActiveSheet
                
                # Get the row number (Excel is 1-indexed)
                row_number = contact_data.get('rowNumber')
                if not row_number:
                    print("Warning: No rowNumber in contact_data, cannot revert source file changes.")
                    return
                
                # Get the row object
                row = worksheet.Rows(row_number)
                
                # Clear Color (xlNone = -4142)
                row.Interior.ColorIndex = -4142  # xlNone (No Color)
                
                # Find and clear the stamp column
                stamp_column = None
                max_col = worksheet.UsedRange.Columns.Count if worksheet.UsedRange else 0
                
                # Search existing columns for _DISPO_STAMP
                for col_idx in range(1, max_col + 1):
                    cell_value = worksheet.Cells(1, col_idx).Value
                    if cell_value and str(cell_value).strip().upper() == "_DISPO_STAMP":
                        stamp_column = col_idx
                        break
                
                # Clear the stamp value if found
                if stamp_column:
                    worksheet.Cells(row_number, stamp_column).Value = ""
                
                # Save the workbook
                workbook.Save()
                print("Reverted source file changes (live editing).")
                return
            
        except Exception as e:
            print(f"Live editing revert failed ({e}), trying fallback...")
        
        # --- ATTEMPT 2: FALLBACK TO CLOSED-FILE EDITING (OPENPYXL) ---
        try:
            if not self.workbook_path:
                print("Warning: No workbook path available for revert.")
                return
            
            workbook = load_workbook(self.workbook_path)
            sheet_name = contact_data.get('sheetName', 'Sheet1')
            
            if sheet_name not in workbook.sheetnames:
                print(f"Warning: Sheet '{sheet_name}' not found in workbook.")
                return
            
            sheet = workbook[sheet_name]
            row_number = contact_data.get('rowNumber')
            if not row_number:
                print("Warning: No rowNumber in contact_data, cannot revert.")
                return
            
            # Find stamp and color columns
            stamp_col = None
            color_col = None
            for col in range(1, sheet.max_column + 1):
                header = sheet.cell(row=1, column=col).value
                if header == "_DISPO_STAMP" or header == "Disposition Stamp":
                    stamp_col = col
                elif header == "Status Color":
                    color_col = col
            
            # Clear the row color (set to white/no fill)
            from openpyxl.styles import PatternFill
            white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            for col_idx in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=row_number, column=col_idx)
                cell.fill = white_fill
            
            # Clear stamp if found
            if stamp_col:
                sheet.cell(row=row_number, column=stamp_col).value = ""
            
            workbook.save(self.workbook_path)
            print("Reverted source file changes (closed file editing).")
            
        except Exception as e:
            print(f"Error reverting source file: {e}")
            import traceback
            traceback.print_exc()
    
    def delete_phone_number_from_source(self, contact_data: Dict[str, Any], phone_number: str):
        """
        Deletes a specific phone number from the source Excel file.
        Searches through all phone columns and removes the matching number.
        
        Args:
            contact_data: Dictionary containing contact data (must include 'rowNumber' and 'sheetName')
            phone_number: The phone number to delete
        """
        # Re-acquire workbook in this thread
        workbook = self._get_workbook_in_thread()
        if not workbook:
            print("Warning: Could not access active workbook to delete phone number.")
            return
        
        try:
            # Get the worksheet name
            sheet_name = contact_data.get('sheetName')
            if not sheet_name or not isinstance(sheet_name, str) or len(sheet_name.strip()) == 0:
                print(f"Warning: Invalid sheetName '{sheet_name}', using active sheet.")
                sheet_name = None
            else:
                sheet_name = str(sheet_name).strip()
            
            # Get the worksheet
            if sheet_name:
                try:
                    worksheet = workbook.Worksheets(sheet_name)
                except Exception as e:
                    print(f"Warning: Could not access sheet '{sheet_name}': {e}. Using active sheet.")
                    worksheet = workbook.ActiveSheet
            else:
                worksheet = workbook.ActiveSheet
            
            # Get the row number (Excel is 1-indexed)
            row_number = contact_data.get('rowNumber')
            if not row_number:
                print("Warning: No rowNumber in contact_data, cannot delete phone number.")
                return
            
            # Normalize the phone number for comparison (remove spaces, dashes, parentheses)
            import re
            phone_normalized = re.sub(r'[\s\-\(\)]', '', str(phone_number))
            
            # Find phone columns by searching header row (row 1)
            # Common phone column names
            phone_keywords = ['phone', 'tel', 'mobile', 'cell', 'number']
            phone_columns = []
            
            max_col = worksheet.UsedRange.Columns.Count if worksheet.UsedRange else 0
            for col_idx in range(1, max_col + 1):
                header_value = worksheet.Cells(1, col_idx).Value
                if header_value:
                    header_str = str(header_value).strip().lower()
                    if any(keyword in header_str for keyword in phone_keywords):
                        phone_columns.append(col_idx)
            
            if not phone_columns:
                print("Warning: No phone columns found in source file.")
                return
            
            # Search through phone columns and remove the matching number
            found_and_removed = False
            for col_idx in phone_columns:
                cell_value = worksheet.Cells(row_number, col_idx).Value
                if cell_value:
                    cell_str = str(cell_value).strip()
                    cell_normalized = re.sub(r'[\s\-\(\)]', '', cell_str)
                    
                    # Check if this cell contains the phone number
                    if phone_normalized in cell_normalized or cell_normalized in phone_normalized:
                        # Clear the cell
                        worksheet.Cells(row_number, col_idx).Value = ""
                        found_and_removed = True
                        print(f"Removed phone number '{phone_number}' from column {col_idx} (row {row_number})")
            
            if found_and_removed:
                # Save the workbook
                workbook.Save()
                print(f"Successfully deleted phone number '{phone_number}' from source file.")
            else:
                print(f"Warning: Phone number '{phone_number}' not found in source file row {row_number}.")
        
        except Exception as e:
            print(f"Error deleting phone number from source file: {e}")
            import traceback
            traceback.print_exc()
